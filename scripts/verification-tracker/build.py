#!/usr/bin/env python3
"""Build the content verification tracker workbook for the review team.

The workbook lists every piece of site text a scholar, mentor or professional
should verify, one text section per row, with the English and Bengali text,
the review it needs, the sources it cites, and blank columns for the team's
verdicts and comments.

Usage:
    python3 scripts/verification-tracker/build.py [--out docs/content-verification-tracker.xlsx]

Requires python3 and openpyxl (pip install openpyxl). Reads locales/en and
locales/bn, plus the item classifications under
scripts/verification-tracker/classification/*.json (review type, priority,
suggested reviewer and key points per item). Items without a classification
fall back to keyword rules.
"""

from __future__ import annotations

import argparse
import datetime as dt
import glob
import json
import os
import re
import subprocess
import sys
from collections import Counter, OrderedDict

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
EN_DIR = os.path.join(ROOT, "locales", "en")
BN_DIR = os.path.join(ROOT, "locales", "bn")
CLASS_DIR = os.path.join(HERE, "classification")
SITE_URL = "https://www.revertguide.com"
CELL_LIMIT = 32000

# ---------------------------------------------------------------------------
# Review taxonomy
# ---------------------------------------------------------------------------

REVIEW_TYPES = [
    "Fiqh ruling",
    "Aqidah (belief)",
    "Quran or hadith citation",
    "Seerah or history",
    "Arabic text or transliteration",
    "Mental health",
    "Suicide-safe language",
    "Safety or abuse",
    "Legal (Canada)",
    "Financial (zakat, riba, money)",
    "Medical or health",
    "Family, marriage or intimacy",
    "Canadian services or contacts",
    "External resource or teacher",
    "Tone and pastoral wording",
    "Site policy or legal notice",
]

REVIEW_TYPE_HELP = {
    "Fiqh ruling": "What is obligatory, permitted or forbidden, and how worship is performed.",
    "Aqidah (belief)": "Statements about Allah, prophets, revelation, the afterlife and creed.",
    "Quran or hadith citation": "A verse or hadith is quoted, paraphrased or cited: check reference, wording and grading.",
    "Seerah or history": "Claims about the Prophet's life, the companions or Islamic history.",
    "Arabic text or transliteration": "Arabic script, transliteration or translation of a phrase, dua or term.",
    "Mental health": "Emotional wellbeing, anxiety, depression, therapy, grief, waswas.",
    "Suicide-safe language": "Suicide or self-harm is mentioned: check wording and that 911 comes before 9-8-8.",
    "Safety or abuse": "Domestic violence, coercion, personal safety, emergencies.",
    "Legal (Canada)": "Law, rights, wills, marriage law, immigration, workplace or school accommodation.",
    "Financial (zakat, riba, money)": "Zakat rules or numbers, interest, mortgages, banking, debt.",
    "Medical or health": "Menstruation, illness and fasting, medication, pregnancy, disability, hygiene.",
    "Family, marriage or intimacy": "Marriage, spouses, divorce, parenting, intimacy, telling family.",
    "Canadian services or contacts": "Phone numbers, organisations, services, dates or prices that must stay current.",
    "External resource or teacher": "A named teacher, channel, book, app or organisation: confirm it is mainstream Sunni and suitable.",
    "Tone and pastoral wording": "Sensitive wording for a new Muslim: encouraging and non-judgemental.",
    "Site policy or legal notice": "About, privacy, terms, accessibility and source policy statements.",
}

SENSITIVE_TYPES = {
    "Fiqh ruling",
    "Aqidah (belief)",
    "Suicide-safe language",
    "Safety or abuse",
    "Legal (Canada)",
    "Financial (zakat, riba, money)",
    "Medical or health",
    "Mental health",
}

ROW_SPECIFIC_TYPES = {
    "Suicide-safe language",
    "Canadian services or contacts",
    "Arabic text or transliteration",
    "Quran or hadith citation",
}

RELIGIOUS_TYPES = {
    "Fiqh ruling",
    "Aqidah (belief)",
    "Quran or hadith citation",
    "Seerah or history",
    "Arabic text or transliteration",
}

# Row-level keyword rules. They only ADD types to what the item classification says.
ROW_RULES = [
    ("Suicide-safe language", re.compile(r"suicid|self-harm|self harm|kill (?:my|your|him|her|them)sel|end(?:ing)? (?:my|your|their|his|her) (?:own )?life|hurt(?:ing)? (?:my|your)self|take (?:my|your|their) (?:own )?life", re.I)),
    ("Quran or hadith citation", re.compile(r"\b(?:Bukhari|Sahih Muslim|Muslim \d|Tirmidhi|Abu Dawud|Abu Dawood|Nasa'?i|Ibn Majah|Muwatta|Musnad|Riyad|hadith|ahadith|Qur'?an \d|Quran \d|Surah|Surat|Sura\b|ayah|verse \d)", re.I)),
    ("Arabic text or transliteration", re.compile(r"[؀-ۿ]")),
    ("Canadian services or contacts", re.compile(r"\b(?:9-1-1|911|9-8-8|988|1-8\d\d[-\s]\d{3}[-\s]\d{4}|\(\d{3}\) ?\d{3}-\d{4}|\d{3}-\d{3}-\d{4}|Kids Help Phone|Talk Suicide|CAMH|Service Canada|Canada Revenue|Legal Aid|helpline|hotline|crisis line|Hope for Wellness|Wellness Together|Naseeha|Khalil Center|Nisa Helpline)", re.I)),
    ("Legal (Canada)", re.compile(r"\b(?:lawyer|legal(?:ly)?\b|the law\b|court|custody|\bwills\b|\b(?:a|your|my|their) will\b|last will|will and testament|inheritance|estate\b|notar|immigration|refugee|citizenship|human rights|discriminat|employment standards|tenant|landlord|power of attorney|executor|marriage licen[cs]e|legally married|civil marriage)", re.I)),
    ("Financial (zakat, riba, money)", re.compile(r"\b(?:zakat|nisab|riba|interest(?:-| )(?:free|based|bearing)|mortgage|loan|debt|bank(?:ing)?\b|credit card|RRSP|TFSA|taxes?\b|halal invest|insurance|pension|sadaqah|fidya|kaffarah|expiation)", re.I)),
    ("Medical or health", re.compile(r"\b(?:menstruat|menses|period(?:s)?\b|postnatal|postpartum|pregnan|breastfeed|medication|medicine|doctor|physician|nurse|illness|diabet|hospital|surgery|vaccin|disabilit|injur|wound|bleeding|urin|nosebleed|incontinence|contracept|IVF|autopsy|organ donation|palliative|MAID\b|life support|dementia|chronic)", re.I)),
    ("Mental health", re.compile(r"\b(?:anxiety|anxious|depress|therap|counsell?|psycholog|psychiatr|panic|trauma|PTSD|OCD|waswas|burnout|lonel|grie[fv]|bereave|mental health|self-care|overwhelm|intrusive thought|distress)", re.I)),
    ("Safety or abuse", re.compile(r"\b(?:abus(?:e|ive|ed)|violen|assault|harass|threat|unsafe|coerc|stalk|shelter|police|emergency|in danger|forced marriage|honou?r-based)", re.I)),
    ("Family, marriage or intimacy", re.compile(r"\b(?:marri|spouse|husband|wife|divorc|talaq|khula|nikah|mahr|intima|\bsex(?:ual)?\b|dating|boyfriend|girlfriend|in-laws|\bwali\b|polygam|custody)", re.I)),
    ("External resource or teacher", re.compile(r"https?://|\bYouTube\b|\bchannel\b|\bpodcast\b|Yaqeen|SeekersGuidance|Bayyinah|Islamwise|Green Lane|New Muslim Academy|WhyIslam|Sunnah\.com|Quran\.com", re.I)),
]

REVIEWERS = ["Scholar", "Mentor", "Mental health professional", "Legal professional", "Medical professional"]
STATUS_LABEL = {
    "draft": "draft",
    "source-checked": "source-checked",
    "review-needed": "review-needed",
    "approved": "approved",
    None: "not tracked",
    "": "not tracked",
}
STATUS_HELP = OrderedDict([
    ("review-needed", "Touches rulings, mental health, safety, legal, financial or sensitive family matters and still needs a qualified reviewer."),
    ("source-checked", "The cited sources were verified against their links; the content itself still awaits a reviewer."),
    ("draft", "Not yet checked at all."),
    ("approved", "A qualified reviewer approved it."),
    ("not tracked", "This part of the site has no review status field yet (stages, glossary entries, site page text)."),
])

# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

BENGALI_RE = re.compile(r"[ঀ-৿]")
ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def load_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def load_pair(rel):
    en = load_json(os.path.join(EN_DIR, rel))
    bn_path = os.path.join(BN_DIR, rel)
    bn = load_json(bn_path) if os.path.exists(bn_path) else None
    return en, bn


def by_id(items):
    return {x.get("id"): x for x in (items or []) if isinstance(x, dict)}


def g(obj, *path):
    """Nested get that tolerates None, missing keys and short lists."""
    cur = obj
    for p in path:
        if cur is None:
            return None
        if isinstance(p, int):
            if isinstance(cur, list) and 0 <= p < len(cur):
                cur = cur[p]
            else:
                return None
        else:
            if isinstance(cur, dict):
                cur = cur.get(p)
            else:
                return None
    return cur


def s(v):
    return v.strip() if isinstance(v, str) else ""


def bullets(items):
    return "\n".join(f"• {x}" for x in (items or []) if isinstance(x, str) and x.strip())


def paragraphs(items):
    return "\n\n".join(x for x in (items or []) if isinstance(x, str) and x.strip())


def lines(*parts):
    return "\n".join(p for p in parts if p)


def labelled(label, val):
    val = s(val)
    return f"{label}: {val}" if val else ""


def humanize(key):
    key = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", key).replace("_", " ").replace("-", " ")
    return key[:1].upper() + key[1:].lower()


def clean_text(text):
    text = ILLEGAL_RE.sub("", text or "")
    if text.startswith(("=", "+", "-", "@")) and len(text) > 1 and not text.startswith("- "):
        text = " " + text
    if len(text) > CELL_LIMIT:
        text = text[: CELL_LIMIT - 40] + "\n[... truncated for Excel cell limit]"
    return text


def bn_status(en_text, bn_text):
    if not s(bn_text):
        return "Missing"
    if BENGALI_RE.search(bn_text):
        return "Translated"
    if s(bn_text) == s(en_text):
        if re.search(r"[A-Za-z]{3,}", en_text):
            return "Same as English"
        return "Not needed"
    return "Not translated"


def git_commit():
    try:
        return subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, text=True).strip()
    except Exception:  # noqa: BLE001
        return "unknown"


# ---------------------------------------------------------------------------
# Model
# ---------------------------------------------------------------------------

class Formula(str):
    """A cell value that must be written as a formula, not as text."""


class Item:
    def __init__(self, key, area, title, url, status, source_ids, json_file):
        self.key = key
        self.area = area
        self.title = title
        self.url = url
        self.status = status
        self.source_ids = [x for x in (source_ids or []) if isinstance(x, str)]
        self.json_file = json_file
        self.rows = []
        self.cls = None
        # Optional anchors: top-level key -> "#fragment" (generic walker, depth 0)
        self.anchor_rules = {}
        # When not None, a dict with an "id" gets the anchor "#<id><suffix>"
        self.anchor_from_ids = None


ITEMS: list[Item] = []
CLASSIFICATION: dict = {}
SOURCES_BY_ID: dict = {}
SOURCE_USE = Counter()


def new_item(*args, **kwargs):
    it = Item(*args, **kwargs)
    ITEMS.append(it)
    return it


def add_row(item, section, en_text, bn_text, json_path, source_ids=None, anchor=None):
    en_text = s(en_text)
    if not en_text:
        return
    item.rows.append({
        "section": section,
        "en": en_text,
        "bn": s(bn_text),
        "json_path": json_path,
        "source_ids": [x for x in (source_ids or []) if isinstance(x, str)] or None,
        "anchor": anchor or "",
    })


# ---------------------------------------------------------------------------
# Generic walker for heterogeneous JSON (tools, pages, dawah guide, ui.json)
# ---------------------------------------------------------------------------

SKIP_KEYS = {
    "id", "slug", "icon", "src", "videoId", "start", "kind", "category", "stageId", "resourceId",
    "replySource", "coordinatesPrecision", "reviewStatus", "lastSourceChecked", "nodeOrder",
    "startNodeId", "nextNodeId", "ariaLabel", "sourceIds", "referenceIds", "seeAlso", "relatedStepIds",
    "relatedTopicIds", "relatedGlossaryIds", "relatedResourceIds", "resourceIds", "stepIds", "type",
    "units", "chapters", "logoAlt", "searchPlaceholder", "searchLabel", "loading", "error",
    "errorTitle", "liveRegion", "loadingAriaLabel", "listAriaLabel", "latitudePlaceholder",
    "longitudePlaceholder", "locationPlaceholder", "emptyAction", "emptyPrefix", "buttonLabel",
    "linkLabel", "footerLinkLabel", "footerPrefix", "footerSuffix", "seasonalButton", "convertedLinkLabel",
    "helpLabel", "officialLabel", "prompt", "eyebrow", "goodNextQuestions",
}
LINK_KEYS = {"href", "url", "officialHref", "linkHref", "warningLinkHref", "sourceUrl"}
NAV_LIST_KEYS = {"links", "relatedLinks"}
TITLE_KEYS = ["title", "heading", "name", "question", "term", "phrase", "word", "day", "label", "citation"]
NO_LABEL_KEYS = {
    "body", "content", "text", "summary", "intro", "description", "desc", "explanation", "answer",
    "solution", "response", "definition", "paragraphs", "items", "points", "steps", "subSteps", "note",
    "notes", "tips", "help", "value",
}
PARAGRAPH_KEYS = {"paragraphs", "explanation"}
KEY_LABELS = {
    "cannotReciteYet": "If you cannot recite yet", "invalidatesPrayer": "What invalidates the prayer",
    "sujudAlSahw": "Prostration of forgetfulness (sujud al-sahw)", "notOwed": "Prayers from before Islam are not owed",
    "seatedPrayer": "Praying seated or lying down", "moreCards": "More card", "situationIndex": "When something comes up",
    "tashahhudPlacement": "Where the tashahhud sits", "prayerTable": "Prayer times table", "prayerSequence": "Prayer sequence step",
    "recitations": "Recitation", "phraseVideos": "Phrase video", "phraseSeries": "Phrase series", "mainVideo": "Main video",
    "shapeImage": "Illustration: one rak'ah", "wuduImage": "Illustration: wudu steps", "ghuslImage": "Illustration: ghusl steps", "wuduBreaksImage": "Illustration: what breaks wudu",
    "image": "Illustration", "wuduSteps": "Wudu step", "ghuslSteps": "Ghusl step", "wuduBreaks": "What breaks wudu",
    "ghuslNeeded": "When ghusl is needed", "wuduBreaksDiffer": "What breaks wudu: where the schools differ",
    "wuduBreaksNote": "What breaks wudu: note", "convertGhusl": "Ghusl for a new Muslim", "tayammum": "Dry ablution (tayammum)",
    "istinja": "Cleaning after the toilet (istinja)", "doubts": "Doubts and waswas", "commonQuestions": "Common question",
    "commonCorrections": "Common corrections", "readingPaths": "Reading path", "weekPlan": "Week plan", "resourceLinks": "Resource link",
    "learnArabic": "Learn Arabic", "fatihahWords": "Al-Fatihah word by word", "sounds": "Sounds to listen for",
    "translationTips": "Translation tips", "vocabulary": "Vocabulary", "sections": "Section", "entries": "Entry", "phrases": "Everyday phrases",
    "otherWords": "Words from other languages", "firstRamadanSections": "First Ramadan section", "ramadanCareSections": "Ramadan care section",
    "supportGroups": "Support group", "directories": "Directories", "reading": "Further reading", "therapist": "Choosing a therapist",
    "items": "Item", "options": "Options", "references": "Reference", "nodes": "Node", "scripts": "Script", "scenarios": "Scenario",
    "teachings": "Teaching", "mistakes": "Mistakes", "signs": "Signs", "practice": "How to practise",
    "policySections": "Policy section", "approachItems": "Approach item", "reveals": "What each option reveals",
    "featureCards": "Feature card", "collections": "Collection", "chooseChecklist": "How to choose", "warnings": "Warnings",
    "policyNotes": "Policy note", "seasons": "Seasons", "paths": "Path", "readings": "Readings", "stages": "Stage", "plan": "Making the plan work",
    "extraLink": "Extra link", "video": "Video", "fatihahVideo": "Al-Fatihah video", "alphabetVideo": "Alphabet video",
    "contacts": "Contacts", "rows": "Rows", "gorap": "GORAP", "suggestions": "Suggestions", "communityConnect": "Community connect",
    "additionalGuidance": "Additional guidance", "checklist": "Checklist", "conversationTips": "Conversation tips", "avoid": "Avoid",
    "askImam": "Ask an imam", "askQualified": "Ask a qualified teacher", "sourcesNote": "Sources note", "crisisBody": "Crisis note",
    "usingPoints": "Using this page", "compactReference": "Compact reference", "mattersMost": "What matters most", "listening": "Listening",
    "respect": "Handling the Quran with respect", "confusingVerse": "When a verse confuses you", "coverage": "Coverage note",
    "locationHelper": "Location helper", "firstVisit": "First visit", "pronunciationKey": "Pronunciation key", "privacyNote": "Privacy note",
    "placeholderNotice": "Placeholder notice", "whatFor": "What you can ask for", "whoReads": "Who reads requests",
    "validation": "Validation", "faithAndHelp": "Faith and help", "commonExperiences": "Common experiences", "selfCare": "Self-care",
    "seekHelp": "When to seek help", "disclaimer": "Disclaimer", "crossLink": "Cross link",
    "whatIsRamadan": "What is Ramadan", "forNewMuslims": "For new Muslims", "fastingBasics": "Fasting basics", "preparing": "Preparing",
    "duringRamadan": "During Ramadan", "tarawihAndCommunity": "Tarawih and community", "laylatAlQadr": "Laylat al-Qadr", "eidAlFitr": "Eid al-Fitr",
    "checklistAlt": "Checklist illustration text", "mission": "Mission", "nameBody": "Why the name", "reviewStates": "Review states",
    "badgeBody": "Review badge", "commitment": "Commitment", "limitations": "Known limitations",
    "registryIntro": "Registry intro", "howToUseBody": "How to use",
    "spiritualPractices": "Spiritual practices", "practicalTips": "Practical tips",
}


def is_id_key(k):
    return k.endswith("Id") or k.endswith("Ids") or k.lower().endswith("href") or k in SKIP_KEYS


def key_label(k):
    return KEY_LABELS.get(k) or humanize(k)


def first_title(d):
    for k in TITLE_KEYS:
        if isinstance(d.get(k), str) and d[k].strip():
            return d[k].strip()
    return ""


def has_external_link(lst):
    for x in lst:
        if isinstance(x, dict):
            for k in LINK_KEYS:
                if isinstance(x.get(k), str) and x[k].startswith("http"):
                    return True
    return False


def dict_is_small(d):
    if "sourceIds" in d:
        return False
    total = 0
    for k, v in d.items():
        if isinstance(v, (dict, list)):
            if isinstance(v, list) and all(isinstance(x, str) for x in v):
                total += sum(len(x) for x in v)
            else:
                return False
        elif isinstance(v, str):
            total += len(v)
    return total <= 320


def small_dict_line(d, bd):
    """One-line rendering of a small dict, e.g. '• Call: 9-1-1' or '• Yes, they believe — Helper text: ...'."""
    bd = bd if isinstance(bd, dict) else {}
    en_parts, bn_parts = [], []
    title_done = False
    for k, v in d.items():
        if is_id_key(k) and k not in LINK_KEYS:
            continue
        bv = bd.get(k)
        if isinstance(v, str) and v.strip():
            if k in LINK_KEYS:
                if v.startswith("http"):
                    en_parts.append(f"Link: {v}")
                    bn_parts.append(f"Link: {v}")
                continue
            if (not title_done and k in TITLE_KEYS) or k in NO_LABEL_KEYS:
                en_parts.append(v.strip())
                bn_parts.append(bv.strip() if isinstance(bv, str) and bv.strip() else "")
                title_done = title_done or k in TITLE_KEYS
            else:
                en_parts.append(f"{humanize(k)}: {v.strip()}")
                bn_parts.append(f"{humanize(k)}: {bv.strip()}" if isinstance(bv, str) and bv.strip() else "")
        elif isinstance(v, list) and v and all(isinstance(x, str) for x in v):
            en_parts.append("; ".join(x.strip() for x in v if x.strip()))
            bn_parts.append("; ".join(x.strip() for x in bv if isinstance(x, str) and x.strip()) if isinstance(bv, list) else "")
    en_line = " — ".join(p for p in en_parts if p)
    if len(en_parts) == 2 and not any(p.startswith(("Link:",)) for p in en_parts) and en_parts[0].endswith(":"):
        en_line = f"{en_parts[0]} {en_parts[1]}"
        bn_line = f"{bn_parts[0] or en_parts[0]} {bn_parts[1] or en_parts[1]}" if any(bn_parts) else ""
    else:
        bn_line = " — ".join(p for p in bn_parts if p)
    return en_line, bn_line


def compose_dict(en, bn, *, split_lists=False):
    """Compose the scalar, string-list and small-list fields of a dict into one text.

    Returns (en_text, bn_text, list_rows, consumed_keys, title_only). list_rows holds
    (key, en_text, bn_text, source_ids) for string lists that get their own row when
    split_lists is True (the top level of an item)."""
    en_parts, bn_parts, list_rows, consumed = [], [], [], set()
    bn = bn if isinstance(bn, dict) else {}
    extra_sources = {k[: -len("SourceIds")]: v for k, v in en.items() if k.endswith("SourceIds") and isinstance(v, list)}

    def sources_for(key):
        for prefix, ids in extra_sources.items():
            if key == prefix or key.startswith(prefix) or prefix.startswith(key):
                return ids
        return None

    title_done = False
    title_text = ""
    for k, v in en.items():
        if is_id_key(k) and k not in LINK_KEYS:
            continue
        bv = bn.get(k)
        if isinstance(v, str):
            if k in LINK_KEYS:
                if v.startswith("http"):
                    en_parts.append(f"Link: {v}")
                    bn_parts.append(f"Link: {bv if isinstance(bv, str) and bv.startswith('http') else v}")
                    consumed.add(k)
                continue
            if not v.strip():
                continue
            consumed.add(k)
            if not title_done and k in TITLE_KEYS:
                en_parts.append(v.strip())
                title_text = v.strip()
                if isinstance(bv, str) and bv.strip():
                    bn_parts.append(bv.strip())
                title_done = True
            elif k in NO_LABEL_KEYS:
                en_parts.append(v.strip())
                if isinstance(bv, str) and bv.strip():
                    bn_parts.append(bv.strip())
            else:
                en_parts.append(f"{key_label(k)}: {v.strip()}")
                if isinstance(bv, str) and bv.strip():
                    bn_parts.append(f"{key_label(k)}: {bv.strip()}")
        elif isinstance(v, list) and v and all(isinstance(x, str) for x in v):
            joiner = paragraphs if k in PARAGRAPH_KEYS else bullets
            en_l = joiner(v)
            bn_l = joiner(bv) if isinstance(bv, list) else ""
            if not en_l:
                continue
            consumed.add(k)
            label = "" if k in NO_LABEL_KEYS else f"{key_label(k)}:\n"
            if split_lists:
                list_rows.append((k, en_l, bn_l, sources_for(k)))
            else:
                en_parts.append(f"{label}{en_l}")
                if bn_l:
                    bn_parts.append(f"{label}{bn_l}")
        elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
            if k in NAV_LIST_KEYS and not has_external_link(v):
                consumed.add(k)
                continue
            if all(dict_is_small(x) for x in v) and len(v) <= 14 and sum(text_length(x) for x in v) <= 900:
                consumed.add(k)
                en_lines, bn_lines = [], []
                for i, x in enumerate(v):
                    el, bl_ = small_dict_line(x, g(bv, i))
                    if el:
                        en_lines.append(f"• {el}")
                        bn_lines.append(f"• {bl_}" if bl_ else "")
                if en_lines:
                    label = "" if k in NO_LABEL_KEYS else f"{key_label(k)}:\n"
                    en_parts.append(label + "\n".join(en_lines))
                    if any(bn_lines):
                        bn_parts.append(label + "\n".join(x for x in bn_lines if x))
    en_text = "\n".join(en_parts)
    title_only = bool(title_text) and en_text.strip() == title_text
    return en_text, "\n".join(bn_parts), list_rows, consumed, title_only


def singular_label(k):
    lab = key_label(k)
    if k in KEY_LABELS:
        return lab
    if lab.endswith("ies"):
        return lab[:-3] + "y"
    if lab.endswith("s") and not lab.endswith("ss"):
        return lab[:-1]
    return lab


def generic_walk(item, en, bn, label_path, json_path, depth=0, inherited_sources=None, anchor=None):
    if isinstance(en, dict):
        if item.anchor_from_ids is not None and isinstance(en.get("id"), str) and en["id"]:
            anchor = f"#{en['id']}{item.anchor_from_ids}"
        own_sources = en.get("sourceIds") if isinstance(en.get("sourceIds"), list) else inherited_sources
        en_text, bn_text, list_rows, consumed, title_only = compose_dict(en, bn, split_lists=(depth == 0))
        has_children = any(
            isinstance(v, (dict, list)) and k not in consumed and not (is_id_key(k) and k not in LINK_KEYS)
            and not (isinstance(v, list) and all(isinstance(x, str) for x in v))
            for k, v in en.items()
        )
        if en_text.strip() and not (title_only and has_children):
            add_row(item, label_path or "Overview", en_text, bn_text, json_path, own_sources, anchor=anchor)
        for k, en_l, bn_l, src_ids in list_rows:
            row_anchor = item.anchor_rules.get(k, anchor) if depth == 0 else anchor
            add_row(item, f"{label_path} › {key_label(k)}" if label_path else key_label(k), en_l, bn_l,
                    f"{json_path}.{k}" if json_path else k, src_ids or own_sources, anchor=row_anchor)
        for k, v in en.items():
            if k in consumed or (is_id_key(k) and k not in LINK_KEYS):
                continue
            bv = bn.get(k) if isinstance(bn, dict) else None
            sub_label = f"{label_path} › {key_label(k)}" if label_path else key_label(k)
            sub_path = f"{json_path}.{k}" if json_path else k
            sub_anchor = item.anchor_rules.get(k, anchor) if depth == 0 else anchor
            if isinstance(v, dict):
                title = first_title(v)
                lbl = f"{sub_label}: {title}" if title and title != key_label(k) else sub_label
                generic_walk(item, v, bv, lbl, sub_path, depth + 1, own_sources, sub_anchor)
            elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                base = f"{label_path} › {singular_label(k)}" if label_path else singular_label(k)
                for i, x in enumerate(v):
                    title = first_title(x)
                    lbl = f"{base} {i + 1} of {len(v)}" + (f": {title}" if title else "")
                    generic_walk(item, x, g(bv, i), lbl, f"{sub_path}[{i}]", depth + 1, own_sources, sub_anchor)
            elif isinstance(v, list) and v and not all(isinstance(x, str) for x in v):
                for i, x in enumerate(v):
                    if isinstance(x, (dict, list)):
                        generic_walk(item, x, g(bv, i), f"{sub_label} {i + 1}", f"{sub_path}[{i}]", depth + 1, own_sources, sub_anchor)
    elif isinstance(en, list):
        if all(isinstance(x, str) for x in en):
            add_row(item, label_path, bullets(en), bullets(bn) if isinstance(bn, list) else "", json_path, inherited_sources, anchor=anchor)
        else:
            for i, x in enumerate(en):
                generic_walk(item, x, g(bn, i), f"{label_path} {i + 1}", f"{json_path}[{i}]", depth + 1, inherited_sources, anchor)


UI_ANCHORS = {
    ("mentalHealth", "validation"): "#validation", ("mentalHealth", "faithAndHelp"): "#islamic-framing",
    ("mentalHealth", "commonExperiences"): "#common-experiences", ("mentalHealth", "selfCare"): "#self-care",
    ("mentalHealth", "seekHelp"): "#seek-help", ("mentalHealth", "recovery"): "#recovery", ("mentalHealth", "resources"): "#resources",
    ("about", "mission"): "#mission", ("about", "nameBody"): "#site-name", ("about", "approachItems"): "#approach",
    ("about", "whoRunsBody"): "#who-runs", ("about", "reviewIntro"): "#review", ("about", "reviewStates"): "#review",
    ("about", "badgeBody"): "#review", ("about", "reviewerBody"): "#review", ("about", "lastReviewed"): "#review",
    ("about", "disclaimer"): "#disclaimer", ("about", "contact"): "#contact",
    ("getHelp", "whatFor"): "#what-for", ("getHelp", "form"): "#send-request", ("getHelp", "whoReads"): "#who-reads",
    ("getHelp", "privacy"): "#privacy", ("getHelp", "report"): "#report", ("getHelp", "contact"): "#contact",
    ("communityGroups", "channel"): "#channel", ("communityGroups", "groups"): "#groups", ("communityGroups", "reveals"): "#reveals",
    ("communityGroups", "rules"): "#rules", ("communityGroups", "report"): "#report",
    ("accessibility", "commitment"): "#commitment", ("accessibility", "standards"): "#standards", ("accessibility", "features"): "#features",
    ("accessibility", "limitations"): "#limitations", ("accessibility", "feedbackBody"): "#feedback", ("accessibility", "lastUpdated"): "#last-updated",
    ("ramadan", "whatIsRamadan"): "#what-is-ramadan", ("ramadan", "forNewMuslims"): "#for-new-muslims", ("ramadan", "fastingBasics"): "#fasting-basics",
    ("ramadan", "preparing"): "#preparing", ("ramadan", "duringRamadan"): "#during-ramadan", ("ramadan", "tarawihAndCommunity"): "#tarawih",
    ("ramadan", "laylatAlQadr"): "#laylat-al-qadr", ("ramadan", "eidAlFitr"): "#eid", ("ramadan", "planningBody"): "#first-ramadan-plan-heading",
    ("ramadan", "careBody"): "#ramadan-care-heading", ("ramadan", "checklistAlt"): "#first-ramadan-plan-heading",
    ("duaDhikr", "usingPoints"): "#using-reference-heading", ("duaDhikr", "crisisBody"): "#crisis-boundary-heading",
    ("duaDhikr", "compactReference"): "#compact-reference-heading",
    ("quranStarter", "mattersMost"): "#matters-most-heading", ("quranStarter", "translationBody"): "#translation-policy-heading",
    ("quranStarter", "listening"): "#listening-heading", ("quranStarter", "respect"): "#respect-heading",
    ("quranStarter", "confusingVerse"): "#confusing-verse-heading", ("quranStarter", "resourcesBody"): "#resources-heading",
    ("quranStarter", "learnArabicBody"): "#learn-arabic", ("quranStarter", "readingPathsBody"): "#reading-paths-heading",
    ("quranStarter", "weekBody"): "#week-plan-heading", ("quranStarter", "wordByWordBody"): "#learn-arabic",
    ("salahCompanion", "mattersMost"): "#matters-most-heading", ("salahCompanion", "askImam"): "#ask-imam-heading",
    ("salahCompanion", "phrasesVideoBody"): "#phrase-videos-heading", ("salahCompanion", "placementBody"): "#placement-heading",
    ("salahCompanion", "shapeBody"): "#shape-heading", ("salahCompanion", "recitationBody"): "#recitation-heading",
    ("salahCompanion", "invalidatesBody"): "#invalidates-heading", ("salahCompanion", "watchBody"): "#watch-heading",
    ("wuduGhusl", "mattersMost"): "#what-matters-heading", ("wuduGhusl", "wuduBody"): "#wudu-heading",
    ("wuduGhusl", "ghuslBody"): "#ghusl-heading", ("wuduGhusl", "askQualified"): "#qualified-help-heading",
    ("findMasjid", "coverage"): "#masjid-coverage-heading", ("findMasjid", "firstVisit"): "#first-visit-help-heading",
    ("stage", "learningPathIntro"): "#learning-path-heading", ("stage", "finalStageBody"): "#next-heading",
}
UI_ID_ANCHOR_PAGES = {("privacy", "policySections"), ("terms", "sections")}


def load_ui_inventory():
    path = os.path.join(HERE, "ui-inventory.json")
    if not os.path.exists(path):
        return {}
    return load_json(path).get("uiPagesContent") or {}


def text_length(o):
    if isinstance(o, str):
        return len(o)
    if isinstance(o, dict):
        return sum(text_length(v) for k, v in o.items() if not is_id_key(k))
    if isinstance(o, list):
        return sum(text_length(v) for v in o)
    return 0


# ---------------------------------------------------------------------------
# Collection walkers
# ---------------------------------------------------------------------------

def video_rows(item, videos, bn_videos, json_path, label="Video", anchor=None):
    for i, v in enumerate(videos or []):
        if not isinstance(v, dict):
            continue
        bv = g(bn_videos, i) or {}
        vid = v.get("videoId", "")
        en_t = lines(
            s(v.get("title")),
            labelled("Channel", v.get("channel")),
            labelled("Length", v.get("duration")),
            s(v.get("note")),
            f"Link: https://www.youtube.com/watch?v={vid}" if vid else "",
        )
        bn_t = lines(
            s(bv.get("title")),
            labelled("Channel", bv.get("channel")),
            labelled("Length", bv.get("duration")),
            s(bv.get("note")),
        )
        add_row(item, f"{label} {i + 1}: {s(v.get('title'))}", en_t, bn_t, f"{json_path}[{i}]", v.get("sourceIds"), anchor=anchor)


def walk_stages():
    en, bn = load_pair("stages.json")
    bn_map = by_id(bn)
    for st in en:
        b = bn_map.get(st["id"], {})
        it = new_item(f"stage:{st['id']}", "Roadmap stage", st["title"], f"/roadmap/{st['id']}", None, st.get("sourceIds"), "locales/en/stages.json")
        base = f"[id={st['id']}]"
        add_row(it, "Title, subtitle and description",
                lines(st.get("title"), st.get("subtitle"), labelled("Duration", st.get("duration")), st.get("description")),
                lines(b.get("title"), b.get("subtitle"), labelled("Duration", b.get("duration")), b.get("description")), f"{base}.title/subtitle/description")
        add_row(it, "Main goal", st.get("mainGoal"), b.get("mainGoal"), f"{base}.mainGoal", anchor="#overview-heading")
        add_row(it, "Success looks like", st.get("success"), b.get("success"), f"{base}.success", anchor="#overview-heading")
        add_row(it, "Don't worry about", bullets(st.get("dontWorry")), bullets(b.get("dontWorry")), f"{base}.dontWorry", anchor="#overview-heading")
        add_row(it, "Focus now", bullets(st.get("focusNow")), bullets(b.get("focusNow")), f"{base}.focusNow", anchor="#stage-guidance-heading")
        add_row(it, "Can wait", bullets(st.get("canWait")), bullets(b.get("canWait")), f"{base}.canWait", anchor="#stage-guidance-heading")
        add_row(it, "Ask for help if", bullets(st.get("askHelpIf")), bullets(b.get("askHelpIf")), f"{base}.askHelpIf", anchor="#stage-guidance-heading")
        lp, blp = st.get("learningPath") or {}, b.get("learningPath") or {}
        for key, label in (("main", "Learning path: main resource"), ("optional", "Learning path: optional resource")):
            r, br = lp.get(key) or {}, blp.get(key) or {}
            if r:
                add_row(it, label, lines(r.get("title"), r.get("body"), labelled("Link", r.get("href")) if s(r.get("href")).startswith("http") else ""),
                        lines(br.get("title"), br.get("body")), f"{base}.learningPath.{key}", anchor="#learning-path-heading")
        add_row(it, "Learning path: guardrails", bullets(lp.get("guardrails")), bullets(blp.get("guardrails")), f"{base}.learningPath.guardrails", anchor="#learning-path-heading")


def walk_steps():
    for path in sorted(glob.glob(os.path.join(EN_DIR, "steps", "*.json"))):
        rel = os.path.relpath(path, EN_DIR)
        st, b = load_pair(rel)
        b = b or {}
        slug = st.get("slug") or st["id"]
        it = new_item(f"step:{st['id']}", "Roadmap step", st["title"], f"/roadmap/{st.get('stageId')}/{slug}", st.get("reviewStatus"), st.get("sourceIds"), f"locales/en/{rel}")
        add_row(it, "Title", lines(st.get("title"), labelled("Time estimate", st.get("timeEstimate"))), lines(b.get("title"), labelled("Time estimate", b.get("timeEstimate"))), "title, timeEstimate")
        add_row(it, "Why it matters", st.get("whyMatters"), b.get("whyMatters"), "whyMatters", anchor="#why-heading")
        acts = st.get("exactActions") or []
        for i, a in enumerate(acts):
            ba = g(b, "exactActions", i) or {}
            add_row(it, f"Action {i + 1} of {len(acts)}", lines(a.get("text"), bullets(a.get("subSteps"))), lines(ba.get("text"), bullets(ba.get("subSteps"))), f"exactActions[{i}]", anchor="#actions-heading")
        obs = st.get("obstacles") or []
        for i, o in enumerate(obs):
            bo = g(b, "obstacles", i) or {}
            add_row(it, f"Obstacle {i + 1} of {len(obs)}", lines(labelled("Problem", o.get("problem")), labelled("Solution", o.get("solution"))),
                    lines(labelled("Problem", bo.get("problem")), labelled("Solution", bo.get("solution"))), f"obstacles[{i}]", anchor="#obstacles-heading")
        add_row(it, "Tiny version", st.get("tinyVersion"), b.get("tinyVersion"), "tinyVersion", anchor="#obstacles-heading")
        add_row(it, "Unlocks next", st.get("unlocksNext"), b.get("unlocksNext"), "unlocksNext", anchor="#unlocks-heading")
        cq = st.get("commonQuestions") or []
        for i, o in enumerate(cq):
            bo = g(b, "commonQuestions", i) or {}
            add_row(it, f"Common question {i + 1} of {len(cq)}", lines(labelled("Q", o.get("problem")), labelled("A", o.get("solution"))),
                    lines(labelled("Q", bo.get("problem")), labelled("A", bo.get("solution"))), f"commonQuestions[{i}]", anchor="#common-questions-heading")
        for key, label, anchor in (("contextNotes", "Context note", "#context-notes-heading"), ("gentleScripts", "Gentle script", "#scripts-heading")):
            for i, n in enumerate(st.get(key) or []):
                bn_n = g(b, key, i) or {}
                add_row(it, f"{label}: {s(n.get('title'))}", lines(n.get("title"), n.get("body")), lines(bn_n.get("title"), bn_n.get("body")), f"{key}[{i}]", anchor=anchor)
        video_rows(it, st.get("videos"), b.get("videos"), "videos", anchor="#videos-heading")


def walk_topics():
    for path in sorted(glob.glob(os.path.join(EN_DIR, "topics", "*.json"))):
        rel = os.path.relpath(path, EN_DIR)
        t, b = load_pair(rel)
        b = b or {}
        slug = t.get("slug") or t["id"]
        it = new_item(f"topic:{t['id']}", "Topic", t["title"], f"/topics/{slug}", t.get("reviewStatus"), t.get("sourceIds"), f"locales/en/{rel}")
        add_row(it, "Title and description", lines(t.get("title"), t.get("description")), lines(b.get("title"), b.get("description")), "title, description")
        secs = t.get("sections") or []
        for i, sec in enumerate(secs):
            bs = g(b, "sections", i) or {}
            add_row(it, f"Section {i + 1} of {len(secs)}: {s(sec.get('heading'))}", lines(sec.get("heading"), sec.get("content")), lines(bs.get("heading"), bs.get("content")), f"sections[{i}]", anchor=f"#section-{i}")
        img, bimg = t.get("image") or {}, b.get("image") or {}
        if img:
            add_row(it, "Illustration text", lines(labelled("Alt text", img.get("alt")), labelled("Caption", img.get("caption")), labelled("Credit", img.get("credit"))),
                    lines(labelled("Alt text", bimg.get("alt")), labelled("Caption", bimg.get("caption")), labelled("Credit", bimg.get("credit"))), "image")
        cards = t.get("linkCards") or []
        if cards:
            en_c = bullets([lines(c.get("label"), c.get("body")).replace("\n", " — ") for c in cards])
            bn_c = bullets([lines(c.get("label"), c.get("body")).replace("\n", " — ") for c in (b.get("linkCards") or []) if isinstance(c, dict)])
            add_row(it, "Link cards", en_c, bn_c, "linkCards", anchor="#link-cards-heading")
        video_rows(it, t.get("videos"), b.get("videos"), "videos", anchor="#topic-videos-heading")


def walk_guides(rel, key_prefix, area, url_prefix):
    en, bn = load_pair(rel)
    bn_map = by_id(bn)
    for gd in en:
        b = bn_map.get(gd["id"], {})
        it = new_item(f"{key_prefix}:{gd['id']}", area, gd["title"], f"{url_prefix}/{gd.get('slug') or gd['id']}", gd.get("reviewStatus"), gd.get("sourceIds"), f"locales/en/{rel}")
        base = f"[id={gd['id']}]"
        add_row(it, "Title, description and intro", lines(gd.get("title"), gd.get("badge"), gd.get("description"), gd.get("intro")),
                lines(b.get("title"), b.get("badge"), b.get("description"), b.get("intro")), f"{base}.title/description/intro")
        for key, label, anchor in (("summary", "Summary", "#start-here-heading"), ("focusNow", "Focus now", "#focus-now-heading"), ("canWait", "Can wait", "#can-wait-heading")):
            add_row(it, label, bullets(gd.get(key)), bullets(b.get(key)), f"{base}.{key}", anchor=anchor)
        secs = gd.get("sections") or []
        for i, sec in enumerate(secs):
            bs = g(b, "sections", i) or {}
            add_row(it, f"Section {i + 1} of {len(secs)}: {s(sec.get('heading'))}",
                    lines(sec.get("heading"), sec.get("body"), bullets(sec.get("items"))),
                    lines(bs.get("heading"), bs.get("body"), bullets(bs.get("items"))), f"{base}.sections[{i}]", anchor=f"#section-{i}")
        for i, sc in enumerate(gd.get("scripts") or []):
            bs = g(b, "scripts", i) or {}
            add_row(it, f"Script: {s(sc.get('title'))}", lines(sc.get("title"), sc.get("body")), lines(bs.get("title"), bs.get("body")), f"{base}.scripts[{i}]", anchor="#scripts-heading")
        for i, sc in enumerate(gd.get("scenarios") or []):
            bs = g(b, "scenarios", i) or {}
            add_row(it, f"Scenario: {s(sc.get('title'))}", lines(sc.get("title"), sc.get("response")), lines(bs.get("title"), bs.get("response")), f"{base}.scenarios[{i}]", anchor="#scenarios-heading")
        video_rows(it, gd.get("videos"), b.get("videos"), f"{base}.videos", anchor="#videos-heading")


def walk_faq():
    en, bn = load_pair("faq.json")
    bn_map = by_id(bn)
    for f in en:
        b = bn_map.get(f["id"], {})
        it = new_item(f"faq:{f['id']}", "FAQ", f.get("question", f["id"]), f"/faq#{f['id']}", f.get("reviewStatus"), f.get("sourceIds"), "locales/en/faq.json")
        rm, brm = f.get("readMore") or {}, b.get("readMore") or {}
        add_row(it, f"Question and answer ({f.get('category', '')})",
                lines(labelled("Q", f.get("question")), labelled("A", f.get("answer")), labelled("Where opinions differ", f.get("differences")), labelled("Read more", lines(rm.get("label"), rm.get("href")).replace("\n", " → "))),
                lines(labelled("Q", b.get("question")), labelled("A", b.get("answer")), labelled("Where opinions differ", b.get("differences")), labelled("Read more", brm.get("label"))),
                f"[id={f['id']}]")


def walk_glossary():
    en, bn = load_pair("glossary.json")
    bn_map = by_id(bn)
    for e in en:
        b = bn_map.get(e["id"], {})
        it = new_item(f"glossary:{e['id']}", "Glossary", e.get("term", e["id"]), f"/glossary#{e['id']}", None, e.get("sourceIds"), "locales/en/glossary.json")
        add_row(it, "Term and definition",
                lines(labelled("Term", e.get("term")), labelled("Arabic", e.get("arabicText")), labelled("Transliteration", e.get("transliteration")), labelled("Definition", e.get("definition")), labelled("See also", ", ".join(e.get("seeAlso") or []))),
                lines(labelled("Term", b.get("term")), labelled("Arabic", b.get("arabicText") or e.get("arabicText")), labelled("Transliteration", b.get("transliteration") or e.get("transliteration")), labelled("Definition", b.get("definition"))),
                f"[id={e['id']}]")


def walk_resources():
    en, bn = load_pair("resources.json")
    bn_map = by_id(bn)
    for r in en:
        b = bn_map.get(r["id"], {})
        it = new_item(f"resource:{r['id']}", "Resource", r.get("title", r["id"]), "/resources", r.get("reviewStatus"), r.get("sourceIds"), "locales/en/resources.json")
        add_row(it, f"Resource ({r.get('type', '')})",
                lines(r.get("title"), labelled("Organisation", r.get("organization")), labelled("Link", r.get("url")), r.get("description"), labelled("Best for", r.get("bestFor")), labelled("Trust note", r.get("trustNote"))),
                lines(b.get("title"), labelled("Organisation", b.get("organization")), b.get("description"), labelled("Best for", b.get("bestFor")), labelled("Trust note", b.get("trustNote"))),
                f"[id={r['id']}]")


TOOL_META = {
    "dua-dhikr": ("Supplication and remembrance (dua and dhikr) tool", "/dua-dhikr"),
    "quran-starter": ("Quran starter tool", "/quran-starter"),
    "salah-companion": ("Learn to pray (salah) tool", "/tools/salah-companion"),
    "wudu-ghusl": ("Washing before prayer (wudu and ghusl) tool", "/tools/wudu-ghusl"),
}
# Page anchors per top-level key (verified against the route files).
TOOL_ANCHORS = {
    "salah-companion": ({
        "cannotReciteYet": "#cannot-recite-heading", "invalidatesPrayer": "#invalidates-heading", "prayerSequence": "#shape-heading",
        "recitations": "#recitation-heading", "commonCorrections": "#corrections-heading", "mainVideo": "#watch-heading",
        "phraseVideos": "#phrase-videos-heading", "phraseSeries": "#phrase-videos-heading", "shapeImage": "#shape-heading",
        "tashahhudPlacement": "#placement-heading", "sujudAlSahw": "#sujud-al-sahw", "prayerTable": "#table-heading",
        "seatedPrayer": "#seated-prayer", "situationIndex": "#when-something-comes-up",
    }, ""),
    "wudu-ghusl": ({
        "wuduSteps": "#wudu-heading", "ghuslSteps": "#ghusl-heading", "wuduBreaks": "#wudu-breaks-heading",
        "wuduBreaksDiffer": "#wudu-breaks-heading", "wuduBreaksNote": "#wudu-breaks-heading", "wuduBreaksImage": "#wudu-breaks-heading",
        "ghuslNeeded": "#ghusl-needed-heading", "commonCorrections": "#corrections-heading", "wuduVideo": "#wudu-heading",
        "ghuslVideo": "#ghusl-heading", "convertGhusl": "#convert-ghusl-heading", "wuduImage": "#wudu-heading", "ghuslImage": "#ghusl-heading",
        "tayammum": "#tayammum-heading", "commonQuestions": "#questions-heading", "doubts": "#doubts",
    }, ""),
    "dua-dhikr": ({}, "-heading"),
    "quran-starter": ({
        "vocabulary": "#vocabulary-heading", "translationTips": "#translation-policy-heading", "readingPaths": "#reading-paths-heading",
        "weekPlan": "#week-plan-heading", "resourceLinks": "#resources-heading", "learnArabic": "#learn-arabic",
    }, None),
}


def walk_tools():
    for stem, (title, url) in TOOL_META.items():
        rel = f"tools/{stem}.json"
        en, bn = load_pair(rel)
        it = new_item(f"tool:{stem}", "Tool page", title, url, en.get("reviewStatus"), None, f"locales/en/{rel}")
        rules, id_suffix = TOOL_ANCHORS.get(stem, ({}, None))
        it.anchor_rules = rules
        it.anchor_from_ids = id_suffix
        generic_walk(it, en, bn or {}, "", "")


def walk_ramadan_page():
    en, bn = load_pair("pages/ramadan-planning.json")
    it = new_item("page:ramadan-planning", "Ramadan guide page", "Ramadan guide: planning sections", "/ramadan", None, None, "locales/en/pages/ramadan-planning.json")
    it.anchor_rules = {"firstRamadanSections": "#first-ramadan-plan-heading", "ramadanCareSections": "#ramadan-care-heading"}
    generic_walk(it, en, bn or {}, "", "")


def walk_dawah():
    en, bn = load_pair("dawah-guides/personal.json")
    bn = bn or {}
    it = new_item("dawah:personal", "Dawah guide (hidden)", en.get("title", "Dawah guide"), f"/{en.get('slug', 'dawah-guide-personal')}", en.get("reviewStatus"), None, "locales/en/dawah-guides/personal.json")
    add_row(it, "Title and audience note", lines(en.get("title"), en.get("subtitle"), en.get("audienceNote"), labelled("Review note", en.get("reviewNote"))),
            lines(bn.get("title"), bn.get("subtitle"), bn.get("audienceNote")), "title/subtitle/audienceNote")
    add_row(it, "Tone principles", bullets(en.get("tonePrinciples")), bullets(bn.get("tonePrinciples")), "tonePrinciples")
    nodes, bnodes = en.get("nodes") or {}, bn.get("nodes") or {}
    order = en.get("nodeOrder") or list(nodes.keys())
    for nid in order:
        n = nodes.get(nid)
        if not isinstance(n, dict):
            continue
        generic_walk(it, n, bnodes.get(nid), f"Node: {s(n.get('title')) or nid}", f"nodes.{nid}", depth=1)
    for key in ("gorap", "suggestions", "communityConnect", "additionalGuidance"):
        if key in en:
            generic_walk(it, en[key], bn.get(key), humanize(key), key, depth=1)
    refs = en.get("references") or []
    for i, r in enumerate(refs):
        br = g(bn, "references", i) or {}
        add_row(it, f"Reference {i + 1} of {len(refs)}: {s(r.get('citation')) or r.get('id', '')}",
                lines(r.get("citation"), r.get("label"), r.get("arabic"), r.get("translation"), labelled("Translator", r.get("translator")), labelled("Source", lines(r.get("sourceName"), r.get("sourceUrl")).replace("\n", " ")), labelled("Review note", r.get("reviewNote")), labelled("Status", r.get("reviewStatus"))),
                lines(br.get("citation"), br.get("label"), br.get("arabic"), br.get("translation"), labelled("Translator", br.get("translator"))),
                f"references[{i}]", r.get("sourceIds"))


def walk_small_collections():
    specs = [
        ("resource-collections.json", "collection:resource-collections", "Resources page", "Resources page: collections, cards, checklist and warnings", "/resources"),
        ("seasonal-calendar.json", "collection:seasonal-calendar", "Seasonal calendar", "Seasonal calendar (dates and notes)", "/seasonal"),
        ("source-categories.json", "collection:source-categories", "Sources page", "Sources page: categories and source policy notes", "/sources"),
    ]
    for rel, key, area, title, url in specs:
        en, bn = load_pair(rel)
        src = en.get("sourceIds") if isinstance(en, dict) else None
        it = new_item(key, area, title, url, None, src, f"locales/en/{rel}")
        generic_walk(it, en, bn or {}, "", "")
    en, bn = load_pair("events.json")
    if en.get("events"):
        it = new_item("collection:events", "Help & community pages", "Events list", "/events", None, None, "locales/en/events.json")
        generic_walk(it, en, bn or {}, "", "")


UI_PAGE_AREA = {
    "mentalHealth": "Mental health page",
    "ramadan": "Ramadan guide page",
    "getHelp": "Help & community pages", "communityGroups": "Help & community pages", "events": "Help & community pages",
    "faq": "Help & community pages", "dawahGuide": "Help & community pages",
    "about": "Site information pages", "privacy": "Site information pages", "terms": "Site information pages",
    "accessibility": "Site information pages", "sources": "Site information pages", "sourceNotes": "Site information pages",
    "home": "Site information pages", "footer": "Site information pages",
    "prayerTimes": "Tool page text", "qibla": "Tool page text", "asmaAlHusna": "Tool page text", "duaDhikr": "Tool page text",
    "quranStarter": "Tool page text", "salahCompanion": "Tool page text", "wuduGhusl": "Tool page text",
}
UI_PAGE_LABEL = {
    "mentalHealth": "Mental health page", "ramadan": "Ramadan guide page", "getHelp": "Get help page", "communityGroups": "Community groups page",
    "events": "Events page", "faq": "FAQ page text", "dawahGuide": "Dawah guide page text", "about": "About page", "privacy": "Privacy policy",
    "terms": "Terms of use", "accessibility": "Accessibility statement", "sources": "Sources page text", "sourceNotes": "Source notes",
    "home": "Home page", "footer": "Footer", "prayerTimes": "Prayer times tool", "qibla": "Qibla tool", "asmaAlHusna": "99 names page",
    "duaDhikr": "Dua and dhikr page text", "quranStarter": "Quran starter page text", "salahCompanion": "Salah companion page text",
    "wuduGhusl": "Wudu and ghusl page text", "guides": "Guides hub", "seasonal": "Seasonal hub", "findMasjid": "Masjid finder",
    "glossary": "Glossary page text", "resourcesPage": "Resources page text", "roadmap": "Roadmap page text", "stage": "Stage page text",
    "step": "Step page text", "topics": "Topics hub", "topic": "Topic page text", "video": "Video player text", "offline": "Offline page",
    "notFound": "404 page", "rootRedirect": "Root redirect",
}
UI_PAGE_URL = {
    "mentalHealth": "/mental-health", "ramadan": "/ramadan", "getHelp": "/get-help", "communityGroups": "/community-groups", "events": "/events",
    "faq": "/faq", "dawahGuide": "/dawah-guide-personal", "about": "/about", "privacy": "/privacy", "terms": "/terms", "accessibility": "/accessibility",
    "sources": "/sources", "sourceNotes": "/sources", "home": "/", "footer": "/", "prayerTimes": "/prayer-times", "qibla": "/qibla",
    "asmaAlHusna": "/asma-al-husna", "duaDhikr": "/dua-dhikr", "quranStarter": "/quran-starter", "salahCompanion": "/tools/salah-companion",
    "wuduGhusl": "/tools/wudu-ghusl", "guides": "/guides", "seasonal": "/seasonal", "findMasjid": "/resources/find-masjid", "glossary": "/glossary",
    "resourcesPage": "/resources", "roadmap": "/roadmap", "stage": "/roadmap", "step": "/roadmap", "topics": "/topics", "topic": "/topics",
    "video": "/", "offline": "/offline", "notFound": "/", "rootRedirect": "/",
}


def walk_ui():
    en, bn = load_pair("ui.json")
    bn = bn or {}
    inventory = load_ui_inventory()
    pages, bpages = en.get("pages") or {}, bn.get("pages") or {}
    entries = []
    for page, sub in pages.items():
        if not isinstance(sub, dict):
            continue
        for key, val in sub.items():
            entries.append((f"ui:pages.{page}.{key}", page, key, val, g(bpages, page, key)))
    entries.append(("ui:footer", "footer", None, en.get("footer") or {}, bn.get("footer") or {}))
    for item_key, page, key, val, bval in entries:
        cls = CLASSIFICATION.get(item_key)
        inv = inventory.get(item_key[3:]) or {}
        if cls is not None:
            include = bool(cls.get("include", True)) or bool(inv)
        elif inv:
            include = True
        else:
            include = text_length(val) >= 200
        if not include:
            continue
        area = UI_PAGE_AREA.get(page, "Other page text")
        page_label = UI_PAGE_LABEL.get(page, humanize(page))
        cls_label = (cls or {}).get("label") if (cls or {}).get("include") else None
        label = cls_label or inv.get("label") or (cls or {}).get("label") or (f"{page_label} > {humanize(key)}" if key else page_label)
        url = (cls or {}).get("url") or UI_PAGE_URL.get(page, "/")
        anchor = UI_ANCHORS.get((page, key))
        inv_url = inv.get("urlEn") or ""
        if inv_url.startswith("/en"):
            inv_url = inv_url[3:] or "/"
            if "#" in inv_url:
                inv_url, _, inv_anchor = inv_url.partition("#")
                anchor = anchor or f"#{inv_anchor}"
            if not (cls or {}).get("url") and inv_url and "{" not in inv_url and "," not in inv_url:
                url = inv_url
        json_path = f"pages.{page}.{key}" if key else page
        it = new_item(item_key, area, label, url, None, None, "locales/en/ui.json")
        if (page, key) in UI_ID_ANCHOR_PAGES:
            it.anchor_from_ids = ""
        if isinstance(val, str):
            add_row(it, humanize(key), val, bval if isinstance(bval, str) else "", json_path, anchor=anchor)
        else:
            generic_walk(it, val, bval, singular_label(key) if isinstance(val, list) and key else "", json_path, anchor=anchor)


SVG_USAGE = {
    "al-fatihah-card": ("Al-Fatihah learning card", "/quran-starter#learn-arabic"),
    "articles-of-faith": ("What Muslims believe (iman)", "/topics/beliefs"),
    "first-ramadan-checklist": ("My first Ramadan checklist", "/ramadan#first-ramadan-plan-heading"),
    "five-pillars": ("The five pillars of Islam", "/topics/five-pillars"),
    "ghusl-steps": ("Full-body wash (ghusl) in five steps", "/tools/wudu-ghusl#ghusl-heading"),
    "prayer-day-timeline": ("The five daily prayers across the day", "/tools/salah-companion#table-heading"),
    "rakah-map-tashahhud": ("Rak'ah map and tashahhud", "/tools/salah-companion#placement-heading"),
    "rakah-steps": ("One unit of prayer (rak'ah), step by step", "/tools/salah-companion#shape-heading"),
    "what-breaks-wudu": ("What breaks wudu", "/tools/wudu-ghusl#wudu-breaks-heading"),
    "wudu-steps": ("Washing before prayer (wudu) in eight steps", "/tools/wudu-ghusl#wudu-heading"),
}


def walk_svgs():
    """Illustrations under public/graphics carry English text that shows in every locale."""
    import xml.etree.ElementTree as ET
    ns = "{http://www.w3.org/2000/svg}"
    for path in sorted(glob.glob(os.path.join(ROOT, "public", "graphics", "*.svg"))):
        stem = os.path.splitext(os.path.basename(path))[0]
        try:
            root = ET.parse(path).getroot()
        except ET.ParseError:
            continue
        title_el, desc_el = root.find(f"{ns}title"), root.find(f"{ns}desc")
        title = s(title_el.text if title_el is not None else "") or SVG_USAGE.get(stem, (stem, "/"))[0]
        desc = s(desc_el.text if desc_el is not None else "")
        texts = []
        for el in root.iter(f"{ns}text"):
            t = re.sub(r"\s+", " ", " ".join(el.itertext())).strip()
            if t and t not in texts:
                texts.append(t)
        url = SVG_USAGE.get(stem, ("", "/"))[1]
        it = new_item(f"svg:{stem}", "Illustration text (SVG)", title, url, None, None, f"public/graphics/{stem}.svg")
        add_row(it, "Title and description (accessible text)", lines(title, desc), "", "<title>, <desc>")
        add_row(it, "Text drawn in the image", "\n".join(texts), "", "<text> elements")


def load_site_notes():
    path = os.path.join(HERE, "site-notes.json")
    return load_json(path) if os.path.exists(path) else []


# ---------------------------------------------------------------------------
# Classification and formatting of rows
# ---------------------------------------------------------------------------

def load_classification():
    merged = {}
    for path in sorted(glob.glob(os.path.join(CLASS_DIR, "*.json"))):
        data = load_json(path)
        for k, v in (data.get("items") or {}).items():
            merged[k] = v
    return merged


def norm_types(types):
    out = []
    for t in types or []:
        if t in REVIEW_TYPES and t not in out:
            out.append(t)
        else:
            low = t.lower()
            match = next((r for r in REVIEW_TYPES if r.lower().split(" (")[0] in low or low in r.lower()), None)
            if match and match not in out:
                out.append(match)
    return out


def row_types_from_text(text):
    found = []
    for name, rx in ROW_RULES:
        if rx.search(text):
            found.append(name)
    return found


def fallback_item_types(item):
    types = set()
    for r in item.rows:
        types.update(row_types_from_text(r["en"]))
    if item.area in ("FAQ", "Tool page", "Roadmap step", "Topic", "Seasonal guide", "Ramadan guide page"):
        types.add("Fiqh ruling")
    if item.area == "Resource":
        types.add("External resource or teacher")
    if item.area == "Illustration text (SVG)":
        # Diagrams of worship: keep the religious types only; the keyword rules add the rest per row.
        text = "\n".join(r["en"] for r in item.rows)
        types = {"Fiqh ruling"}
        if re.search(r"[\u0600-\u06FF]", text):
            types.add("Arabic text or transliteration")
        if ROW_RULES[1][1].search(text):
            types.add("Quran or hadith citation")
    if item.area in ("Site information pages",):
        types.add("Site policy or legal notice")
    if not types:
        types.add("Tone and pastoral wording")
    return [t for t in REVIEW_TYPES if t in types]


def fallback_reviewers(types):
    out = []
    if RELIGIOUS_TYPES & set(types):
        out.append("Scholar")
    if {"Tone and pastoral wording", "Canadian services or contacts", "External resource or teacher", "Family, marriage or intimacy", "Safety or abuse"} & set(types):
        out.append("Mentor")
    if {"Mental health", "Suicide-safe language"} & set(types):
        out.append("Mental health professional")
    if "Legal (Canada)" in types:
        out.append("Legal professional")
    if "Medical or health" in types:
        out.append("Medical professional")
    return out or ["Mentor"]


def priority_for(status, types):
    unreviewed = status in (None, "", "draft", "review-needed")
    if unreviewed and SENSITIVE_TYPES & set(types):
        return "High"
    if unreviewed or (RELIGIOUS_TYPES | {"External resource or teacher"}) & set(types):
        return "Medium"
    return "Low"


def finalize_items():
    for it in ITEMS:
        c = CLASSIFICATION.get(it.key)
        it.cls = c
        if c:
            it.types = norm_types(c.get("reviewTypes")) or fallback_item_types(it)
            it.reviewers = [r for r in (c.get("suggestedReviewers") or []) if r in REVIEWERS] or fallback_reviewers(it.types)
            it.priority = c.get("priority") if c.get("priority") in ("High", "Medium", "Low") else priority_for(it.status, it.types)
            it.key_points = [s(x) for x in (c.get("keyPoints") or []) if s(x)]
            it.flags = [s(x) for x in (c.get("preCheckFlags") or []) if s(x)]
        else:
            it.types = fallback_item_types(it)
            it.reviewers = fallback_reviewers(it.types)
            it.priority = priority_for(it.status, it.types)
            it.key_points = []
            it.flags = []
        for r in it.rows:
            found = set(row_types_from_text(r["en"]))
            # Types that describe a specific passage stay on the rows where the passage is
            # (or on a single-row item); the broader types apply to every row of the item.
            base = [t for t in it.types if t not in ROW_SPECIFIC_TYPES or t in found or len(it.rows) == 1]
            extra = [t for t in found if t not in base]
            r["types"] = base + extra
            r["priority"] = it.priority
            if "Suicide-safe language" in r["types"] or ("Safety or abuse" in extra and it.status != "approved"):
                r["priority"] = "High"
            r["bn_status"] = bn_status(r["en"], r["bn"])
            ids = r["source_ids"] or it.source_ids
            r["sources_n"] = len(ids)
            r["sources_short"] = source_lines(ids, with_url=False)


def source_lines(ids, with_url):
    out = []
    for i in ids:
        src = SOURCES_BY_ID.get(i)
        if not src:
            out.append(f"{i} — (not found in sources.json)")
            continue
        org = f" ({src['organization']})" if s(src.get("organization")) else ""
        line = f"{i} — {src.get('title', '')}{org}"
        if with_url and s(src.get("url")):
            line += f" — {src['url']}"
        out.append(line)
    return "\n".join(out)


def count_source_use():
    def walk(o):
        if isinstance(o, dict):
            for k, v in o.items():
                if k == "sourceIds" and isinstance(v, list):
                    for i in v:
                        if isinstance(i, str):
                            SOURCE_USE[i] += 1
                else:
                    walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
    for path in glob.glob(os.path.join(EN_DIR, "**", "*.json"), recursive=True):
        if path.endswith("sources.json"):
            continue
        walk(load_json(path))


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
INPUT_HEADER_FILL = PatternFill("solid", fgColor="B8860B")
INPUT_FILL = PatternFill("solid", fgColor="FFFBEA")
GEN_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
NOWRAP = Alignment(wrap_text=False, vertical="top")
FILL_HIGH = PatternFill("solid", fgColor="F8D7DA")
FILL_MED = PatternFill("solid", fgColor="FFF3CD")
FILL_LOW = PatternFill("solid", fgColor="E2F0D9")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_AMBER = PatternFill("solid", fgColor="FFEB9C")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_DONE = PatternFill("solid", fgColor="EEF7EE")
FILL_FLAG = PatternFill("solid", fgColor="FCE4D6")

VERDICTS = '"Approved,Needs changes,Rejected,Unsure - ask another reviewer"'
YES_NO = '"Yes,No"'


def style_header(ws, headers, widths, input_cols=(), wrap_cols=None, nowrap_cols=()):
    ws.append(headers)
    for idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = INPUT_HEADER_FILL if h in input_cols else HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
    ws.row_dimensions[1].height = 34


def write_rows(ws, rows, input_cols_idx=(), nowrap_idx=(), link_idx=None, first_data_row=2):
    for r_i, row in enumerate(rows, start=first_data_row):
        for c_i, val in enumerate(row, start=1):
            if isinstance(val, Formula):
                cell = ws.cell(row=r_i, column=c_i, value=str(val))
            else:
                cell = ws.cell(row=r_i, column=c_i, value=clean_text(val) if isinstance(val, str) else val)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = NOWRAP if c_i in nowrap_idx else WRAP
            cell.border = BORDER
            if c_i in input_cols_idx:
                cell.fill = INPUT_FILL
            if link_idx and c_i == link_idx and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.font = Font(name=FONT, size=10, color="0563C1", underline="single")


def add_list_validation(ws, col_letter, formula, last_row):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")


def add_verdict_formats(ws, col_letter, last_row):
    rng = f"{col_letter}2:{col_letter}{last_row}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col_letter}2="Approved"'], fill=FILL_GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col_letter}2="Needs changes"'], fill=FILL_AMBER))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col_letter}2="Rejected"'], fill=FILL_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${col_letter}2,6)="Unsure"'], fill=FILL_AMBER))


REVIEW_HEADERS = [
    "Row ID", "Area", "Item", "Section", "Priority", "Review needed", "Suggested reviewer", "Current status (site)",
    "English text", "Bengali text (বাংলা)", "Bengali status", "Sources (n)", "Sources used (id — title)",
    "Key points to check", "Pre-check flags",
    "English verdict", "English reviewer", "English comments",
    "Bengali verdict", "Bengali reviewer", "Bengali comments",
    "Done", "Date done", "Page URL", "Content ID", "JSON location", "Item key",
]
REVIEW_WIDTHS = [8, 18, 30, 30, 9, 30, 20, 13, 70, 70, 13, 8, 42, 55, 42, 15, 16, 40, 15, 16, 40, 8, 12, 42, 24, 46, 24]
REVIEW_INPUT = ["English verdict", "English reviewer", "English comments", "Bengali verdict", "Bengali reviewer", "Bengali comments", "Done", "Date done"]


def page_url(path, anchor=None):
    path = path or "/"
    base, _, frag = path.partition("#")
    frag = anchor.lstrip("#") if anchor else frag
    url = SITE_URL + "/en" + (base if base != "/" else "")
    return f"{url}#{frag}" if frag else url


def build_review_sheet(wb):
    ws = wb.create_sheet("Review rows")
    style_header(ws, REVIEW_HEADERS, REVIEW_WIDTHS, REVIEW_INPUT)
    col = {h: i + 1 for i, h in enumerate(REVIEW_HEADERS)}
    rows = []
    n = 0
    for it in ITEMS:
        content_id = it.key.split(":", 1)[1]
        for r in it.rows:
            n += 1
            rows.append([
                f"R{n:04d}", it.area, it.title, r["section"], r["priority"], ", ".join(r["types"]), ", ".join(it.reviewers),
                STATUS_LABEL.get(it.status, it.status or "not tracked"),
                r["en"], r["bn"], r["bn_status"], r["sources_n"], r["sources_short"],
                "\n".join(f"• {k}" for k in it.key_points), "\n".join(f"⚠ {f}" for f in it.flags),
                None, None, None, None, None, None, None, None,
                page_url(it.url, r.get("anchor")), content_id, f"{it.json_file} › {r['json_path']}", it.key,
            ])
    input_idx = tuple(col[h] for h in REVIEW_INPUT)
    write_rows(ws, rows, input_idx, nowrap_idx=(col["Sources used (id — title)"],), link_idx=col["Page URL"])
    last = len(rows) + 1
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_HEADERS))}{last}"
    L = {h: get_column_letter(i) for h, i in col.items()}
    for h in ("English verdict", "Bengali verdict"):
        add_list_validation(ws, L[h], VERDICTS, last)
        add_verdict_formats(ws, L[h], last)
    add_list_validation(ws, L["Done"], YES_NO, last)
    add_list_validation(ws, L["Bengali status"], '"Translated,Missing,Same as English,Not translated,Not needed"', last)
    p = L["Priority"]
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="High"'], fill=FILL_HIGH))
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="Medium"'], fill=FILL_MED))
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="Low"'], fill=FILL_LOW))
    b = L["Bengali status"]
    ws.conditional_formatting.add(f"{b}2:{b}{last}", FormulaRule(formula=[f'AND(${b}2<>"Translated",${b}2<>"Not needed")'], fill=FILL_AMBER))
    f = L["Pre-check flags"]
    ws.conditional_formatting.add(f"{f}2:{f}{last}", FormulaRule(formula=[f'LEN(${f}2)>0'], fill=FILL_FLAG))
    d = L["Done"]
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(REVIEW_HEADERS))}{last}", FormulaRule(formula=[f'${d}2="Yes"'], fill=FILL_DONE))
    ws.sheet_properties.tabColor = "1F4E5F"
    return len(rows), L


ITEM_HEADERS = [
    "Item key", "Area", "Item", "Priority", "Review needed", "Suggested reviewer", "Current status (site)",
    "Key points to check", "Pre-check flags", "Sources (n)", "Sources used (id — title — link)",
    "Text rows", "Rows done", "Rows approved (EN)", "Rows approved (BN)", "Progress",
    "Assigned reviewer(s)", "Item verdict", "Item notes", "Page URL", "JSON file",
]
ITEM_WIDTHS = [26, 18, 40, 9, 32, 22, 13, 60, 42, 8, 60, 8, 8, 10, 10, 12, 20, 15, 40, 44, 34]
ITEM_INPUT = ["Assigned reviewer(s)", "Item verdict", "Item notes"]


def build_items_sheet(wb, review_cols, n_review_rows):
    ws = wb.create_sheet("Items")
    R = n_review_rows + 1
    style_header(ws, ITEM_HEADERS, ITEM_WIDTHS, ITEM_INPUT)
    col = {h: i + 1 for i, h in enumerate(ITEM_HEADERS)}
    rk, rdone, ren, rbn = review_cols["Item key"], review_cols["Done"], review_cols["English verdict"], review_cols["Bengali verdict"]
    rows = []
    for i, it in enumerate(ITEMS, start=2):
        rows.append([
            it.key, it.area, it.title, it.priority, ", ".join(it.types), ", ".join(it.reviewers),
            STATUS_LABEL.get(it.status, it.status or "not tracked"),
            "\n".join(f"• {k}" for k in it.key_points), "\n".join(f"⚠ {f}" for f in it.flags),
            len(it.source_ids), source_lines(it.source_ids, with_url=True),
            Formula(f"=COUNTIF('Review rows'!${rk}$2:${rk}${R},$A{i})"),
            Formula(f"=COUNTIFS('Review rows'!${rk}$2:${rk}${R},$A{i},'Review rows'!${rdone}$2:${rdone}${R},\"Yes\")"),
            Formula(f"=COUNTIFS('Review rows'!${rk}$2:${rk}${R},$A{i},'Review rows'!${ren}$2:${ren}${R},\"Approved\")"),
            Formula(f"=COUNTIFS('Review rows'!${rk}$2:${rk}${R},$A{i},'Review rows'!${rbn}$2:${rbn}${R},\"Approved\")"),
            Formula(f"=IF(L{i}=0,\"\",IF(M{i}=L{i},\"Done\",IF(M{i}>0,\"In progress\",\"Not started\")))"),
            None, None, None,
            page_url(it.url), it.json_file,
        ])
    write_rows(ws, rows, tuple(col[h] for h in ITEM_INPUT), nowrap_idx=(col["Sources used (id — title — link)"],), link_idx=col["Page URL"])
    last = len(rows) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ITEM_HEADERS))}{last}"
    L = {h: get_column_letter(i) for h, i in col.items()}
    add_list_validation(ws, L["Item verdict"], VERDICTS, last)
    add_verdict_formats(ws, L["Item verdict"], last)
    p = L["Priority"]
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="High"'], fill=FILL_HIGH))
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="Medium"'], fill=FILL_MED))
    ws.conditional_formatting.add(f"{p}2:{p}{last}", FormulaRule(formula=[f'${p}2="Low"'], fill=FILL_LOW))
    pr = L["Progress"]
    ws.conditional_formatting.add(f"{pr}2:{pr}{last}", FormulaRule(formula=[f'${pr}2="Done"'], fill=FILL_GREEN))
    ws.conditional_formatting.add(f"{pr}2:{pr}{last}", FormulaRule(formula=[f'${pr}2="In progress"'], fill=FILL_AMBER))
    f = L["Pre-check flags"]
    ws.conditional_formatting.add(f"{f}2:{f}{last}", FormulaRule(formula=[f'LEN(${f}2)>0'], fill=FILL_FLAG))
    for r in range(2, last + 1):
        for h in ("Text rows", "Rows done", "Rows approved (EN)", "Rows approved (BN)", "Progress"):
            ws.cell(row=r, column=col[h]).fill = GEN_FILL
    return L


def build_sources_sheet(wb, sources, bn_sources):
    ws = wb.create_sheet("Sources")
    headers = ["Source ID", "Title", "Organisation", "Category", "Source type", "Label", "URL", "Accessed", "Current status (site)",
               "Note (what it supports)", "Note (Bengali)", "Used by (references)", "Verdict", "Reviewer", "Comments"]
    widths = [30, 40, 26, 18, 24, 16, 46, 11, 14, 60, 60, 10, 15, 16, 40]
    inputs = ["Verdict", "Reviewer", "Comments"]
    style_header(ws, headers, widths, inputs)
    col = {h: i + 1 for i, h in enumerate(headers)}
    bn_map = by_id(bn_sources)
    rows = []
    for src in sources:
        b = bn_map.get(src.get("id"), {})
        rows.append([
            src.get("id"), src.get("title"), src.get("organization"), src.get("category"), src.get("sourceType"), src.get("label"),
            src.get("url"), src.get("accessed"), src.get("reviewStatus"), src.get("note"), b.get("note"), SOURCE_USE.get(src.get("id"), 0),
            None, None, None,
        ])
    write_rows(ws, rows, tuple(col[h] for h in inputs), nowrap_idx=(col["URL"],), link_idx=col["URL"])
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    L = {h: get_column_letter(i) for h, i in col.items()}
    add_list_validation(ws, L["Verdict"], VERDICTS, last)
    add_verdict_formats(ws, L["Verdict"], last)
    st = L["Current status (site)"]
    ws.conditional_formatting.add(f"{st}2:{st}{last}", FormulaRule(formula=[f'${st}2="review-needed"'], fill=FILL_MED))
    return L


MASJID_NOTE_KEYS = [("notes", "Notes"), ("visitorNotes", "Visitor notes"), ("womenSpaceNote", "Women's space"), ("newMuslimSupportNote", "New Muslim support"),
                    ("accessibilityNote", "Accessibility"), ("jumuahNote", "Jumu'ah"), ("classSupportNote", "Classes"), ("parkingNote", "Parking")]


def build_masjids_sheet(wb, masjids, bn_masjids):
    ws = wb.create_sheet("Masjids")
    headers = ["Masjid ID", "Name", "City", "Province", "Address", "Postal code", "Phone", "Email", "Website", "Services", "Notes (English)",
               "Notes (Bengali)", "Map pin", "Current status (site)", "Sources (id — title)", "Verdict", "Reviewer", "Comments"]
    widths = [30, 34, 14, 14, 30, 10, 15, 26, 36, 24, 60, 60, 11, 14, 40, 15, 16, 40]
    inputs = ["Verdict", "Reviewer", "Comments"]
    style_header(ws, headers, widths, inputs)
    col = {h: i + 1 for i, h in enumerate(headers)}
    bn_map = by_id(bn_masjids)
    rows = []
    for m in masjids:
        b = bn_map.get(m.get("id"), {})
        en_notes = lines(*[labelled(lbl, m.get(k)) for k, lbl in MASJID_NOTE_KEYS])
        bn_notes = lines(*[labelled(lbl, b.get(k)) for k, lbl in MASJID_NOTE_KEYS])
        pin = m.get("coordinatesPrecision") or ("street" if m.get("coordinates") else "none")
        rows.append([
            m.get("id"), m.get("name"), m.get("city"), m.get("stateProvince"), m.get("address"), m.get("postalCode"), m.get("phone"), m.get("email"),
            m.get("website"), ", ".join(m.get("serviceIds") or []), en_notes, bn_notes, pin, m.get("reviewStatus"),
            source_lines(m.get("sourceIds") or [], with_url=False), None, None, None,
        ])
    write_rows(ws, rows, tuple(col[h] for h in inputs), nowrap_idx=(col["Website"], col["Sources (id — title)"]), link_idx=col["Website"])
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    L = {h: get_column_letter(i) for h, i in col.items()}
    add_list_validation(ws, L["Verdict"], VERDICTS, last)
    add_verdict_formats(ws, L["Verdict"], last)
    st = L["Current status (site)"]
    ws.conditional_formatting.add(f"{st}2:{st}{last}", FormulaRule(formula=[f'${st}2="review-needed"'], fill=FILL_MED))
    return L


def build_notes_sheet(wb, notes):
    ws = wb.create_sheet("Site team notes")
    headers = ["#", "Area", "Finding", "Suggested action", "How it was found", "Status", "Owner", "Comment"]
    widths = [5, 22, 90, 50, 30, 12, 16, 40]
    inputs = ["Status", "Owner", "Comment"]
    style_header(ws, headers, widths, inputs)
    col = {h: i + 1 for i, h in enumerate(headers)}
    rows = [[i, n.get("area"), n.get("note"), n.get("action"), n.get("from"), None, None, None] for i, n in enumerate(notes, start=1)]
    write_rows(ws, rows, tuple(col[h] for h in inputs))
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    L = {h: get_column_letter(i) for h, i in col.items()}
    add_list_validation(ws, L["Status"], '"Open,Fixed,Won\'t fix,Needs scholar"', last)
    st = L["Status"]
    ws.conditional_formatting.add(f"{st}2:{st}{last}", FormulaRule(formula=[f'${st}2="Fixed"'], fill=FILL_GREEN))
    ws.conditional_formatting.add(f"{st}2:{st}{last}", FormulaRule(formula=[f'${st}2="Needs scholar"'], fill=FILL_AMBER))
    return len(rows)


def build_summary_sheet(wb, review_cols, item_cols, src_cols, masjid_cols, n_sources, n_masjids, n_review_rows):
    ws = wb.create_sheet("Summary")
    R = n_review_rows + 1
    I = len(ITEMS) + 1
    ws.column_dimensions["A"].width = 36
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 14
    bold = Font(name=FONT, bold=True, size=10)
    normal = Font(name=FONT, size=10)
    head = Font(name=FONT, bold=True, color="FFFFFF", size=10)

    def header(row, labels):
        for i, lab in enumerate(labels, start=1):
            c = ws.cell(row=row, column=i, value=lab)
            c.font = head
            c.fill = HEADER_FILL
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = BORDER

    def bound(formula):
        formula = re.sub(r"'Review rows'!\$([A-Z]+):\$([A-Z]+)", lambda m: f"'Review rows'!${m.group(1)}$2:${m.group(2)}${R}", formula)
        formula = re.sub(r"Items!\$([A-Z]+):\$([A-Z]+)", lambda m: f"Items!${m.group(1)}$2:${m.group(2)}${I}", formula)
        formula = re.sub(r"Sources!\$([A-Z]+):\$([A-Z]+)", lambda m: f"Sources!${m.group(1)}$2:${m.group(2)}${n_sources + 1}", formula)
        formula = re.sub(r"Masjids!\$([A-Z]+):\$([A-Z]+)", lambda m: f"Masjids!${m.group(1)}$2:${m.group(2)}${n_masjids + 1}", formula)
        return formula

    def put(row, values, font=normal, fmt=None):
        for i, v in enumerate(values, start=1):
            if isinstance(v, str) and v.startswith("="):
                v = bound(v)
            c = ws.cell(row=row, column=i, value=v)
            c.font = font
            c.border = BORDER
            c.alignment = Alignment(vertical="top")
            if fmt and i > 1:
                c.number_format = fmt

    RA, RE, RF, RG, RK, RV, RP, RS = (review_cols[h] for h in ("Area", "Priority", "Review needed", "Suggested reviewer", "Bengali status", "Done", "English verdict", "Bengali verdict"))
    IA = item_cols["Area"]
    ws.cell(row=1, column=1, value="Revert Guide content verification: progress summary").font = Font(name=FONT, bold=True, size=14)
    ws.cell(row=2, column=1, value="All numbers are formulas over the other sheets and update as the team fills in the yellow columns.").font = Font(name=FONT, italic=True, size=10)

    r = 4
    header(r, ["Area", "Items", "Text rows", "High priority rows", "Rows done", "% done", "EN approved", "BN approved"])
    areas = []
    for it in ITEMS:
        if it.area not in areas:
            areas.append(it.area)
    first = r + 1
    for area in areas:
        r += 1
        put(r, [
            area,
            f"=COUNTIF(Items!${IA}:${IA},$A{r})",
            f"=COUNTIF('Review rows'!${RA}:${RA},$A{r})",
            f"=COUNTIFS('Review rows'!${RA}:${RA},$A{r},'Review rows'!${RE}:${RE},\"High\")",
            f"=COUNTIFS('Review rows'!${RA}:${RA},$A{r},'Review rows'!${RV}:${RV},\"Yes\")",
            f"=IF(C{r}=0,0,E{r}/C{r})",
            f"=COUNTIFS('Review rows'!${RA}:${RA},$A{r},'Review rows'!${RP}:${RP},\"Approved\")",
            f"=COUNTIFS('Review rows'!${RA}:${RA},$A{r},'Review rows'!${RS}:${RS},\"Approved\")",
        ])
        ws.cell(row=r, column=6).number_format = "0%"
    r += 1
    last_area = r - 1
    put(r, ["Total", f"=SUM(B{first}:B{last_area})", f"=SUM(C{first}:C{last_area})", f"=SUM(D{first}:D{last_area})", f"=SUM(E{first}:E{last_area})",
            f"=IF(C{r}=0,0,E{r}/C{r})", f"=SUM(G{first}:G{last_area})", f"=SUM(H{first}:H{last_area})"], font=bold)
    ws.cell(row=r, column=6).number_format = "0%"

    r += 2
    header(r, ["Review type", "Text rows", "Rows done", "% done"])
    for t in REVIEW_TYPES:
        r += 1
        put(r, [t, f"=COUNTIF('Review rows'!${RF}:${RF},\"*\"&$A{r}&\"*\")",
                f"=COUNTIFS('Review rows'!${RF}:${RF},\"*\"&$A{r}&\"*\",'Review rows'!${RV}:${RV},\"Yes\")", f"=IF(B{r}=0,0,C{r}/B{r})"])
        ws.cell(row=r, column=4).number_format = "0%"

    r += 2
    header(r, ["Suggested reviewer", "Text rows", "Rows done", "% done"])
    for t in REVIEWERS:
        r += 1
        put(r, [t, f"=COUNTIF('Review rows'!${RG}:${RG},\"*\"&$A{r}&\"*\")",
                f"=COUNTIFS('Review rows'!${RG}:${RG},\"*\"&$A{r}&\"*\",'Review rows'!${RV}:${RV},\"Yes\")", f"=IF(B{r}=0,0,C{r}/B{r})"])
        ws.cell(row=r, column=4).number_format = "0%"

    r += 2
    header(r, ["Priority", "Text rows", "Rows done", "% done"])
    for t in ("High", "Medium", "Low"):
        r += 1
        put(r, [t, f"=COUNTIF('Review rows'!${RE}:${RE},$A{r})", f"=COUNTIFS('Review rows'!${RE}:${RE},$A{r},'Review rows'!${RV}:${RV},\"Yes\")", f"=IF(B{r}=0,0,C{r}/B{r})"])
        ws.cell(row=r, column=4).number_format = "0%"

    r += 2
    header(r, ["Bengali translation status", "Text rows"])
    for t in ("Translated", "Missing", "Same as English", "Not translated", "Not needed"):
        r += 1
        put(r, [t, f"=COUNTIF('Review rows'!${RK}:${RK},$A{r})"])

    r += 2
    header(r, ["Verdicts", "English", "Bengali"])
    for t in ("Approved", "Needs changes", "Rejected", "Unsure - ask another reviewer"):
        r += 1
        put(r, [t, f"=COUNTIF('Review rows'!${RP}:${RP},$A{r})", f"=COUNTIF('Review rows'!${RS}:${RS},$A{r})"])

    r += 2
    header(r, ["Other sheets", "Records", "Marked review-needed", "Verdicts given"])
    SV, SS = src_cols["Verdict"], src_cols["Current status (site)"]
    MV, MS = masjid_cols["Verdict"], masjid_cols["Current status (site)"]
    r += 1
    put(r, ["Sources", n_sources, f"=COUNTIF(Sources!${SS}:${SS},\"review-needed\")", f"=COUNTA(Sources!${SV}2:${SV}{n_sources + 1})"])
    r += 1
    put(r, ["Masjids", n_masjids, f"=COUNTIF(Masjids!${MS}:${MS},\"review-needed\")", f"=COUNTA(Masjids!${MV}2:${MV}{n_masjids + 1})"])
    ws.freeze_panes = "A4"


README_EN = [
    ("h1", "Revert Guide: content verification tracker"),
    ("p", "One workbook for the review team. Every piece of site text that a scholar, mentor or professional should verify is listed on the 'Review rows' sheet, one text section per row, with the English and the Bengali text side by side, the kind of review it needs, the sources it cites, and yellow columns for your verdict and comments."),
    ("h2", "How to use it"),
    ("li", "Open 'Review rows'. Use the filters in the header row to pick your part: filter 'Area' (for example FAQ, Topic, Roadmap step), 'Suggested reviewer' (Scholar, Mentor, Mental health professional, Legal professional, Medical professional) or 'Priority'. Start with High."),
    ("li", "Read the English text. If you review Bengali, read the Bengali text next to it and judge both the meaning and the translation."),
    ("li", "Check the claims against 'Sources used' (the ids point to the 'Sources' sheet, which has every link) and against 'Key points to check'."),
    ("li", "Fill in the yellow columns only: 'English verdict' (Approved, Needs changes, Rejected, Unsure), 'English reviewer' (your name), 'English comments' (what to change and why; quote the exact words). Bengali reviewers use the three Bengali columns the same way."),
    ("li", "Set 'Done' to Yes once the row is settled in both languages. The 'Items' and 'Summary' sheets count progress automatically."),
    ("li", "Grey and white columns are generated from the site's content files: do not edit them. If a text is wrong, say so in the comments and the site team will fix the content file named in 'JSON location'."),
    ("li", "'Sources' and 'Masjids' are separate sheets with their own verdict columns: sources need a citation check (collection, number, grading, working link); masjids need a directory check (mainstream Sunni, details current)."),
    ("li", "'Site team notes' lists cross-cutting findings from the preparation of this workbook (duplicate entries, missing source ids, naming-rule slips, facts to confirm, code-defined text). They are for the site team and a scholar where marked; each has a status column."),
    ("li", "Not covered here: the 99 names of Allah page, whose names and English meanings are fetched from the AlAdhan API at run time rather than stored in the site; and the live prayer-times data. Everything else that a visitor can read is on one of the sheets."),
    ("h2", "What the generated columns mean"),
    ("kv", "Priority", "High: unreviewed text that carries rulings, belief, suicide or self-harm, safety, legal, financial, medical or mental-health guidance. Medium: unreviewed practical text, or source-checked text that still makes religious claims. Low: everything else."),
    ("kv", "Review needed", "The kinds of verification the row needs (see the list below). A row can need several."),
    ("kv", "Suggested reviewer", "Who should look first. Scholars for religious content, mentors for lived-experience and practical guidance, professionals for clinical, legal or medical claims."),
    ("kv", "Current status (site)", "The reviewStatus recorded in the content file today: " + " ".join(f"'{k}' = {v}" for k, v in STATUS_HELP.items())),
    ("kv", "Bengali status", "Translated: Bengali text exists. Missing: no Bengali text for this section. Same as English / Not translated: the Bengali file still carries the English words. Not needed: the text is Arabic, a link or a number and needs no translation."),
    ("kv", "Sources used", "The source ids the site cites for this text, with title and organisation, one per line. The column does not wrap, so click the cell (or read the formula bar) to see the whole list, and look the id up on the 'Sources' sheet for the link and the note on what it supports. The 'Items' sheet repeats the list with links."),
    ("kv", "Key points to check", "A machine-generated reading list of the specific claims, citations, numbers and names in the item. It is a pointer, not a verdict. Empty when nothing specific stands out."),
    ("kv", "Pre-check flags", "Machine-generated warnings about a possible breach of the site's own rules (for example suicide-safe wording, 911 before 9-8-8, a missing source). Please confirm or dismiss them in your comments."),
    ("kv", "Page URL / JSON location", "Where the text appears on the site and where it lives in the content files, for the site team."),
    ("h2", "Review types"),
] + [("kv", t, REVIEW_TYPE_HELP[t]) for t in REVIEW_TYPES] + [
    ("h2", "House rules the content must follow"),
    ("li", "Mainstream Sunni scholarship only. Differences between the four Sunni schools are acknowledged without choosing one. No Shia or Ahmadiyya source or institution."),
    ("li", "Headings lead with plain English and put the Arabic term in brackets, for example 'Washing before prayer (wudu)'. Ghusl is a full-body wash, never a shower."),
    ("li", "Suicide-safe writing: say 'died by suicide' and 'thoughts of suicide'; never 'committed suicide'. Give 911 first for immediate danger, then 9-8-8 (call or text). Never frame suicidal thoughts as a sin. Keep clinical and religious guidance in separate sentences."),
    ("li", "Every religious claim carries source ids that point at sources.json. Quran and hadith links use Quran.com and Sunnah.com. Nothing is invented."),
    ("li", "Tone: encouraging and non-judgemental, written for someone in their first year as a Muslim in Canada."),
]

README_BN = [
    ("h2", "বাংলা পর্যালোচকদের জন্য সংক্ষিপ্ত নির্দেশনা"),
    ("li", "'Review rows' শিটে যান। শিরোনাম সারির ফিল্টার দিয়ে আপনার অংশ বেছে নিন: 'Area' (যেমন FAQ, Topic, Roadmap step), 'Suggested reviewer' (Scholar, Mentor ইত্যাদি) বা 'Priority'। প্রথমে High দিয়ে শুরু করুন।"),
    ("li", "'English text' কলামের ইংরেজি লেখা এবং তার পাশে 'Bengali text (বাংলা)' কলামের বাংলা অনুবাদ পড়ুন। অর্থ ঠিক আছে কি না এবং অনুবাদ সঠিক কি না, দুটোই দেখুন।"),
    ("li", "'Sources used' কলামে দেওয়া সূত্রগুলোর সঙ্গে দাবিগুলো মিলিয়ে দেখুন। প্রতিটি সূত্রের লিংক 'Sources' শিটে আছে। 'Key points to check' কলামে কোন কোন বিষয় যাচাই করতে হবে তার তালিকা আছে।"),
    ("li", "শুধু হলুদ কলামগুলো পূরণ করুন। ইংরেজি লেখার জন্য 'English verdict', 'English reviewer', 'English comments'; বাংলা অনুবাদের জন্য 'Bengali verdict', 'Bengali reviewer', 'Bengali comments'। সিদ্ধান্ত: Approved (অনুমোদিত), Needs changes (পরিবর্তন দরকার), Rejected (বাতিল), Unsure (অন্য পর্যালোচকের মত দরকার)। মন্তব্যে ঠিক কোন শব্দ বদলাতে হবে এবং কেন, তা লিখুন।"),
    ("li", "দুই ভাষার কাজ শেষ হলে 'Done' কলামে Yes দিন। 'Items' ও 'Summary' শিটে অগ্রগতি স্বয়ংক্রিয়ভাবে গণনা হয়।"),
    ("li", "ধূসর ও সাদা কলামগুলো সাইটের কনটেন্ট ফাইল থেকে স্বয়ংক্রিয়ভাবে তৈরি; সেগুলো সম্পাদনা করবেন না। কোনো লেখা ভুল হলে মন্তব্যে লিখুন, সাইট টিম কনটেন্ট ফাইল ঠিক করবে।"),
    ("li", "'Bengali status' কলামে Missing, Same as English বা Not translated থাকলে বুঝবেন সেই অংশের বাংলা অনুবাদ এখনো নেই বা অসম্পূর্ণ।"),
]


def build_readme(wb, meta):
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    ws.sheet_view.showGridLines = False
    r = 1

    def write(kind, a, b=None):
        nonlocal r
        if kind == "h1":
            c = ws.cell(row=r, column=1, value=a)
            c.font = Font(name=FONT, bold=True, size=16, color="1F4E5F")
            r += 1
        elif kind == "h2":
            r += 1
            c = ws.cell(row=r, column=1, value=a)
            c.font = Font(name=FONT, bold=True, size=12, color="1F4E5F")
            r += 1
        elif kind == "p":
            c = ws.cell(row=r, column=1, value=a)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            ws.row_dimensions[r].height = 15 * max(1, len(a) // 150 + 1)
            r += 1
        elif kind == "li":
            ws.cell(row=r, column=1, value="•").font = Font(name=FONT, size=10)
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="right", vertical="top")
            c = ws.cell(row=r, column=2, value=a)
            c.font = Font(name=FONT, size=10)
            c.alignment = WRAP
            ws.row_dimensions[r].height = 14 * max(1, len(a) // 130 + 1)
            r += 1
        elif kind == "kv":
            c = ws.cell(row=r, column=1, value=a)
            c.font = Font(name=FONT, bold=True, size=10)
            c.alignment = WRAP
            d = ws.cell(row=r, column=2, value=b)
            d.font = Font(name=FONT, size=10)
            d.alignment = WRAP
            ws.row_dimensions[r].height = 14 * max(1, len(b) // 130 + 1)
            r += 1

    for entry in README_EN[:2]:
        write(*entry)
    write("kv", "Generated", f"{meta['generated']} from the content files in the repository (last commit {meta['commit']}). Re-run scripts/verification-tracker/build.py to refresh it after content changes; the yellow columns are not carried over, so copy them before regenerating.")
    write("kv", "Contents", f"{meta['rows']} text rows across {meta['items']} items on 'Review rows'; {meta['sources']} sources on 'Sources'; {meta['masjids']} masjid records on 'Masjids'; {meta['notes']} findings on 'Site team notes'. The counts on 'Items' and 'Summary' are formulas and calculate when the file opens (Excel may ask to save on close because of this).")
    write("kv", "Colour key", "Dark header = generated column (read only). Gold header and pale yellow cells = for the review team to fill in. Grey cells = formulas.")
    for entry in README_EN[2:]:
        write(*entry)
    for entry in README_BN:
        write(*entry)
    ws.sheet_properties.tabColor = "B8860B"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=os.path.join(ROOT, "docs", "content-verification-tracker.xlsx"))
    ap.add_argument("--dump-json", help="Also write the generated rows as JSON (for checks)")
    args = ap.parse_args()

    global CLASSIFICATION, SOURCES_BY_ID
    CLASSIFICATION = load_classification()
    sources, bn_sources = load_pair("sources.json")
    SOURCES_BY_ID = by_id(sources)
    count_source_use()

    walk_stages()
    walk_steps()
    walk_topics()
    walk_guides("life-guides.json", "guide", "Life guide", "/guides")
    walk_guides("seasonal-guides.json", "seasonal", "Seasonal guide", "/seasonal")
    walk_faq()
    walk_glossary()
    walk_tools()
    walk_ramadan_page()
    walk_ui()
    walk_dawah()
    walk_resources()
    walk_small_collections()
    walk_svgs()
    finalize_items()

    masjids, bn_masjids = load_pair("masjids.json")

    wb = Workbook()
    n_rows, review_cols = build_review_sheet(wb)
    item_cols = build_items_sheet(wb, review_cols, n_rows)
    src_cols = build_sources_sheet(wb, sources, bn_sources)
    masjid_cols = build_masjids_sheet(wb, masjids, bn_masjids)
    build_summary_sheet(wb, review_cols, item_cols, src_cols, masjid_cols, len(sources), len(masjids), n_rows)
    notes = load_site_notes()
    n_notes = build_notes_sheet(wb, notes)
    meta = {
        "generated": dt.date.today().isoformat(),
        "commit": git_commit(),
        "rows": n_rows,
        "items": len(ITEMS),
        "sources": len(sources),
        "masjids": len(masjids),
        "notes": n_notes,
    }
    build_readme(wb, meta)
    # No cached values are written for formulas; Excel, Google Sheets and LibreOffice calculate them on open.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.move_sheet("Summary", offset=1 - wb.sheetnames.index("Summary"))  # Read me, Summary, Review rows, ...
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)

    missing_cls = [it.key for it in ITEMS if it.cls is None]
    print(f"wrote {args.out}: {n_rows} rows, {len(ITEMS)} items, {len(sources)} sources, {len(masjids)} masjids")
    print(f"items without classification: {len(missing_cls)}")
    if args.dump_json:
        with open(args.dump_json, "w", encoding="utf-8") as fh:
            json.dump([{"key": it.key, "area": it.area, "title": it.title, "url": it.url, "status": it.status, "types": it.types,
                        "priority": it.priority, "rows": it.rows} for it in ITEMS], fh, ensure_ascii=False, indent=1)
    by_area = Counter()
    for it in ITEMS:
        by_area[it.area] += len(it.rows)
    for area, n in by_area.most_common():
        print(f"  {area}: {n} rows")


if __name__ == "__main__":
    sys.exit(main())
